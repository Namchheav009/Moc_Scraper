"""Check exported artifacts, including contact text and spreadsheet formulas."""

from __future__ import annotations

import csv
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from moc_scraper.excel_exporter import COLUMNS, export_results, normalize_date


class NormalizeDateTests(unittest.TestCase):
    def test_supported_site_formats(self):
        for value in (
            "Sep 21, 2026", "September 21, 2026", "21 Sep 2026",
            "21-Sep-26", "21/09/2026", "2026-09-21",
            "2026-09-21T13:14:15", date(2026, 9, 21),
            datetime(2026, 9, 21, 13, 14), "២១/០៩/២០២៦",
        ):
            with self.subTest(value=value):
                self.assertEqual(normalize_date(value), "21-Sep-26")

    def test_unknown_and_empty_dates_are_not_fabricated(self):
        self.assertEqual(normalize_date("not published"), "not published")
        self.assertEqual(normalize_date("31/02/2026"), "31/02/2026")
        self.assertEqual(normalize_date(None), "")
        self.assertEqual(normalize_date("8-Sep-26"), "08-Sep-26")


class ExportResultsTests(unittest.TestCase):
    def test_workbook_and_csv_preserve_data_and_format(self):
        rows = [
            {"No.": 49, "Company Name": "VIVI FOOD CO., LTD.",
             "Phone": "+855 96 770 6842", "Email": "845765518@qq.com",
             "Registration Date": "Sep 8, 2026"},
            {"Company Name": "ក្រុមហ៊ុន", "Phone": "012345678",
             "Email": "info@example.com", "Registration Date": "21/09/2026"},
            {"No.": "0007", "Company Name": "=HYPERLINK(\"https://example.com\")",
             "Phone": "=1+1", "Email": "=1+2", "Registration Date": "=1+3"},
        ]
        with tempfile.TemporaryDirectory() as temporary:
            excel_path, csv_path = export_results(rows, Path(temporary) / "exports")
            self.assertEqual(excel_path.name, "MOC_Businesses.xlsx")
            self.assertEqual(csv_path.name, "MOC_Businesses.csv")
            workbook = load_workbook(excel_path)
            try:
                sheet = workbook.active
                self.assertEqual([cell.value for cell in sheet[1]], COLUMNS)
                self.assertTrue(all(cell.font.bold for cell in sheet[1]))
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:E4")
                self.assertEqual(
                    [sheet.column_dimensions[column].width for column in "ABCDE"],
                    [8, 45, 22, 35, 20],
                )
                self.assertEqual(sheet["A2"].value, 49)
                self.assertEqual(sheet["A3"].value, 2)
                self.assertEqual(sheet["A4"].value, "0007")
                self.assertEqual(sheet["B3"].value, "ក្រុមហ៊ុន")
                for address, expected in {"C2": "+855 96 770 6842", "C3": "012345678", "D2": "845765518@qq.com"}.items():
                    self.assertEqual(sheet[address].value, expected)
                    self.assertEqual(sheet[address].data_type, "s")
                    self.assertEqual(sheet[address].number_format, "@")
                self.assertEqual(sheet["E2"].value, datetime(2026, 9, 8))
                self.assertEqual(sheet["E2"].number_format, "DD-MMM-YY")
                for column in "BCDE":
                    self.assertEqual(sheet[f"{column}4"].data_type, "s")
            finally:
                workbook.close()
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                backup = list(reader)
                self.assertEqual(reader.fieldnames, COLUMNS)
            self.assertEqual(backup[0]["Phone"], "+855 96 770 6842")
            self.assertEqual(backup[1]["Phone"], "012345678")
            self.assertEqual(backup[1]["Company Name"], "ក្រុមហ៊ុន")
            self.assertEqual(backup[0]["Registration Date"], "08-Sep-26")
            self.assertEqual(backup[2]["Company Name"], rows[2]["Company Name"])

    def test_empty_export_has_headers_and_no_placeholder_business(self):
        with tempfile.TemporaryDirectory() as temporary:
            excel_path, csv_path = export_results([], temporary)
            workbook = load_workbook(excel_path)
            try:
                self.assertEqual(workbook.active.max_row, 1)
                self.assertEqual(workbook.active.auto_filter.ref, "A1:E1")
                self.assertEqual([cell.value for cell in workbook.active[1]], COLUMNS)
            finally:
                workbook.close()
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual(list(csv.reader(handle)), [COLUMNS])

    def test_partial_export_keeps_completed_files_and_full_date_year(self):
        with tempfile.TemporaryDirectory() as temporary:
            complete_excel, complete_csv = export_results([], temporary)
            original_excel = complete_excel.read_bytes()
            original_csv = complete_csv.read_bytes()
            partial_excel, partial_csv = export_results(
                [{"Company Name": "Example", "Registration Date": "2069-09-21"}],
                temporary, basename="MOC_Businesses.partial",
            )
            self.assertEqual(partial_excel.name, "MOC_Businesses.partial.xlsx")
            self.assertEqual(partial_csv.name, "MOC_Businesses.partial.csv")
            self.assertEqual(complete_excel.read_bytes(), original_excel)
            self.assertEqual(complete_csv.read_bytes(), original_csv)
            workbook = load_workbook(partial_excel)
            try:
                self.assertEqual(workbook.active["E2"].value, datetime(2069, 9, 21))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
